# Xuất dữ liệu cho trang hub (claude.ai/artifact/72zNCsPigyqbV6P8kQtji5):
#   hub/schedule.json  từ plan/lich-noi-dung-tiktok.xlsx (sheet "Lịch đăng")
#   hub/scripts.json   từ scripts/*/*.json (bỏ kịch bản có "archived": true)
# Cách dùng (trong thư mục video): python3 tools/hub_export.py
# Sau đó publish lại hub với files videos.json, schedule.json, scripts.json.
import glob, json, os
from datetime import datetime
from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
os.chdir(ROOT)

VISUAL = {"serp": "Màn hình Google", "chat": "Màn hình chatbot", "columns": "Biểu đồ cột", "hbars": "Thanh ngang",
          "continue": "Tiếp biểu đồ trước", "waffle": "Ô vuông 100", "list": "Danh sách", "article": "Thẻ bài viết",
          "follow": "Thẻ Follow Adtek", "bignumber": "Số lớn", "versus": "Đối đầu", "slope": "Đường dốc",
          "trend": "Đường xu hướng", "donut": "Vòng tròn", "people": "Hình người", "funnel": "Phễu",
          "quiz": "Đố số liệu (3 lựa chọn)", "myth": "Phá hiểu lầm (đóng dấu)", "shot": "Ảnh chụp bài viết"}

def cell(v):
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    return "" if v is None else str(v)

# Lịch đăng
ws = load_workbook("plan/lich-noi-dung-tiktok.xlsx", data_only=True)["Lịch đăng"]
head = [c.value for c in ws[1]]
key = {"STT": "stt", "Thứ": "weekday", "Ngày đăng": "date", "Dịch vụ": "service", "Chủ đề video": "topic",
       "Câu hỏi mở đầu": "hook", "Số liệu chính": "data", "Nguồn gốc": "sources", "Link nguồn gốc": "links",
       "Kiểm tra nguồn": "check", "Lưu ý khi dùng": "notes", "Từ khóa người xem tìm": "keywords",
       "Nhu cầu tìm kiếm (ước tính)": "demand", "Trạng thái": "status"}
schedule = []
for row in ws.iter_rows(min_row=2, values_only=True):
    if not any(row):
        continue
    item = {key[h]: cell(v) for h, v in zip(head, row) if h in key}
    schedule.append(item)
json.dump(schedule, open("hub/schedule.json", "w"), ensure_ascii=False, indent=1)

# Kịch bản
scripts = []
for path in sorted(glob.glob("scripts/*/*.json")):
    s = json.load(open(path))
    if s.get("archived"):
        continue
    man = f"public/voice/{s['id']}/manifest.json"
    secs = None
    if os.path.exists(man):
        m = json.load(open(man))
        secs = round(sum(x["duration"] + 0.35 for x in m["scenes"]))
    scenes = []
    for sc in s["scenes"]:
        v = sc.get("visual") or {}
        scenes.append({"kicker": sc.get("kicker", ""), "title": (sc.get("headline", "") + " " + sc.get("accent", "")).strip(),
                       "voice": sc["voice"], "visual": VISUAL.get(v.get("type"), "Câu chốt, chữ lớn"),
                       "source": v.get("source", "")})
    first = s["scenes"][0]
    scripts.append({"id": s["id"], "date": s.get("date", ""), "status": s.get("status", ""),
                    "title": (first.get("headline", "") + " " + first.get("accent", "")).strip(),
                    "caption": s.get("caption", ""), "hashtags": s.get("hashtags", []),
                    "words": sum(len(x["voice"].split()) for x in s["scenes"]), "seconds": secs, "scenes": scenes,
                    "review": s.get("review", "")})
scripts.sort(key=lambda x: (x["date"] == "", x["date"]), reverse=False)
json.dump(scripts, open("hub/scripts.json", "w"), ensure_ascii=False, indent=1)
print(f"schedule.json: {len(schedule)} dòng, scripts.json: {len(scripts)} kịch bản")
