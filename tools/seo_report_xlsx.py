"""Xuất báo cáo audit SEO ra Excel (tiếng Việt).

Cách dùng: python3 tools/seo_report_xlsx.py crawl.json issues.json links.json out.xlsx
"""
import json
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="002D72")
SEV_LABEL = {"high": "Cao", "medium": "Trung bình", "low": "Thấp", "info": "Thông tin"}
SEV_ORDER = {"high": 0, "medium": 1, "low": 2, "info": 3}
SEV_FILL = {"Cao": "F8D7DA", "Trung bình": "FFF3CD", "Thấp": "E2E3E5", "Thông tin": "E2E3E5"}
GROUP_LABEL = {"post": "Bài viết", "page": "Trang", "mona_glossary": "Thuật ngữ", "mona_solution": "Giải pháp",
               "mona_reports": "Báo cáo", "mona_recruitment": "Tuyển dụng", "category": "Chuyên mục"}

ISSUES = {
    "status": ("Trang trả về lỗi", "Sửa hoặc gỡ URL khỏi sitemap, đặt chuyển hướng 301 nếu trang đã đổi địa chỉ."),
    "redirect_in_sitemap": ("URL trong sitemap bị chuyển hướng", "Cập nhật sitemap hoặc link về URL đích cuối cùng."),
    "noindex_in_sitemap": ("Trang noindex nhưng vẫn nằm trong sitemap", "Bỏ noindex nếu muốn lên Google, hoặc loại khỏi sitemap."),
    "title_missing": ("Thiếu thẻ title", "Điền SEO title trong Yoast."),
    "title_long": ("Title dài hơn 60 ký tự, Google sẽ cắt bớt", "Rút gọn SEO title trong Yoast, đưa từ khóa chính lên đầu."),
    "title_short": ("Title ngắn hơn 30 ký tự", "Bổ sung từ khóa và lợi ích vào SEO title."),
    "title_duplicate": ("Title trùng với trang khác", "Viết title riêng cho từng trang."),
    "description_missing": ("Thiếu meta description", "Viết meta description 120 đến 155 ký tự trong Yoast, có từ khóa và lời kêu gọi."),
    "description_long": ("Meta description dài hơn 160 ký tự", "Rút gọn còn 120 đến 155 ký tự."),
    "description_short": ("Meta description ngắn hơn 70 ký tự", "Viết dài hơn, nêu rõ nội dung và lợi ích."),
    "description_duplicate": ("Meta description trùng với trang khác", "Viết mô tả riêng cho từng trang."),
    "h1_missing": ("Thiếu thẻ H1", "Thêm 1 tiêu đề H1 chứa từ khóa chính (thường cần sửa template của theme)."),
    "h1_multiple": ("Có nhiều hơn 1 thẻ H1", "Giữ 1 H1, đổi các H1 còn lại thành H2."),
    "canonical_missing": ("Thiếu canonical", "Kiểm tra cài đặt Yoast cho loại nội dung này."),
    "canonical_other": ("Canonical trỏ sang URL khác", "Kiểm tra có chủ đích không, nếu không thì để canonical trỏ về chính trang."),
    "og_image_missing": ("Thiếu ảnh chia sẻ mạng xã hội (og:image)", "Đặt ảnh đại diện (featured image) hoặc ảnh Social trong Yoast."),
    "thin_content": ("Nội dung mỏng (dưới 300 từ)", "Bổ sung nội dung hữu ích, hoặc noindex nếu trang không cần lên Google."),
    "img_alt_missing": ("Ảnh trong nội dung thiếu alt", "Điền Alt Text cho ảnh trong Thư viện, mô tả đúng nội dung ảnh."),
    "slow": ("Thời gian tải HTML trên 3 giây", "Kiểm tra cache WP-Optimize cho trang này, giảm kích thước trang."),
    "crawl_blocked": ("Bị tường lửa chặn khi quét", "Không cần sửa, kiểm tra lại thủ công."),
}


def style_header(ws, headers, widths):
    ws.append(headers)
    for i, w in enumerate(widths, 1):
        c = ws.cell(row=1, column=i)
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"


def finish(ws):
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name=FONT, size=10)
            c.alignment = Alignment(vertical="top", wrap_text=True)
    ws.auto_filter.ref = ws.dimensions


def build(crawl, issues, links, out):
    wb = Workbook()
    summary = wb.active
    summary.title = "Tổng quan"

    # Vấn đề
    ws = wb.create_sheet("Vấn đề")
    style_header(ws, ["Mức độ", "Mã lỗi", "Vấn đề", "Nhóm", "URL", "Chi tiết", "Cách sửa"], [12, 22, 38, 12, 60, 40, 55])
    for i in sorted(issues, key=lambda x: (SEV_ORDER[x["severity"]], x["code"], x["url"])):
        name, fix = ISSUES[i["code"]]
        sev = SEV_LABEL[i["severity"]]
        ws.append([sev, i["code"], name, GROUP_LABEL.get(i["group"], i["group"]), i["url"], i["detail"], fix])
        ws.cell(row=ws.max_row, column=1).fill = PatternFill("solid", fgColor=SEV_FILL[sev])
    finish(ws)
    last_issue = ws.max_row

    # Chi tiết URL
    ws = wb.create_sheet("Chi tiết URL")
    style_header(ws, ["URL", "Nhóm", "HTTP", "Title", "Độ dài title", "Meta description", "Độ dài description",
                      "H1", "Số H1", "Số từ", "Ảnh thiếu alt (nội dung)", "Thời gian tải (giây)", "Canonical", "Robots"],
                 [55, 12, 7, 45, 10, 55, 11, 40, 7, 9, 12, 11, 45, 30])
    alt_detail = {i["url"]: int(i["detail"].split()[0]) for i in issues if i["code"] == "img_alt_missing"}
    for n, r in enumerate(crawl, start=2):
        ws.append([r["url"], GROUP_LABEL.get(r["group"], r["group"]), r["status"], r["title"], f"=LEN(D{n})",
                   r["description"], f"=LEN(F{n})", " | ".join(r["h1"]), len(r["h1"]), r["words"],
                   alt_detail.get(r["url"], 0), r["seconds"], r["canonical"] or "", r["robots"]])
    finish(ws)
    last_url = ws.max_row

    # Link hỏng
    ws = wb.create_sheet("Link hỏng")
    style_header(ws, ["Link", "HTTP", "Số trang chứa link", "Ví dụ trang chứa link"], [60, 8, 12, 60])
    broken = sorted((l for l in links if l["status"] >= 400 or l["status"] == 0), key=lambda l: -len(l["found_on"]))
    for l in broken:
        ws.append([l["link"], l["status"], len(l["found_on"]), "\n".join(l["found_on"][:3])])
    if not broken:
        ws.append(["Không phát hiện link nội bộ hỏng", "", "", ""])
    finish(ws)

    # Tổng quan
    s = summary
    s.column_dimensions["A"].width = 52
    s.column_dimensions["B"].width = 14
    s.column_dimensions["C"].width = 14
    s["A1"] = "AUDIT SEO WEBSITE ADTEK.AGENCY"
    s["A1"].font = Font(name=FONT, bold=True, size=14, color="002D72")
    s["A2"] = "Nguồn: quét tự động toàn bộ URL trong sitemap Yoast (sitemap_index.xml)."
    s["A2"].font = Font(name=FONT, italic=True, size=9)
    rows = [
        ("Số URL đã quét", f"=COUNTA('Chi tiết URL'!A2:A{last_url})"),
        ("Số URL trả về HTTP 200", f"=COUNTIF('Chi tiết URL'!C2:C{last_url},200)"),
        ("Tổng số vấn đề", f"=COUNTA('Vấn đề'!A2:A{last_issue})"),
        ("Vấn đề mức Cao", f"=COUNTIF('Vấn đề'!A2:A{last_issue},\"Cao\")"),
        ("Vấn đề mức Trung bình", f"=COUNTIF('Vấn đề'!A2:A{last_issue},\"Trung bình\")"),
        ("Vấn đề mức Thấp", f"=COUNTIF('Vấn đề'!A2:A{last_issue},\"Thấp\")"),
        ("Link nội bộ hỏng", len(broken)),
    ]
    s["A4"], s["B4"] = "Chỉ số", "Giá trị"
    for c in ("A4", "B4"):
        s[c].font = Font(name=FONT, bold=True, color="FFFFFF")
        s[c].fill = HEAD_FILL
    for i, (label, val) in enumerate(rows, start=5):
        s.cell(row=i, column=1, value=label).font = Font(name=FONT)
        s.cell(row=i, column=2, value=val).font = Font(name=FONT, bold=True)
    s.cell(row=4 + len(rows), column=3, value="Đếm từ sheet Link hỏng").font = Font(name=FONT, italic=True, size=9)

    start = 6 + len(rows)
    s.cell(row=start, column=1, value="Vấn đề").font = Font(name=FONT, bold=True, color="FFFFFF")
    s.cell(row=start, column=2, value="Số URL").font = Font(name=FONT, bold=True, color="FFFFFF")
    s.cell(row=start, column=3, value="Mức độ").font = Font(name=FONT, bold=True, color="FFFFFF")
    for col in (1, 2, 3):
        s.cell(row=start, column=col).fill = HEAD_FILL
    codes = sorted({i["code"] for i in issues}, key=lambda c: (min(SEV_ORDER[i["severity"]] for i in issues if i["code"] == c), c))
    for k, code in enumerate(codes, start=start + 1):
        sev = SEV_LABEL[next(i["severity"] for i in issues if i["code"] == code)]
        s.cell(row=k, column=1, value=ISSUES[code][0]).font = Font(name=FONT)
        s.cell(row=k, column=2, value=f"=COUNTIF('Vấn đề'!B2:B{last_issue},\"{code}\")").font = Font(name=FONT)
        c = s.cell(row=k, column=3, value=sev)
        c.font = Font(name=FONT)
        c.fill = PatternFill("solid", fgColor=SEV_FILL[sev])
    wb.save(out)


if __name__ == "__main__":
    crawl = json.load(open(sys.argv[1], encoding="utf-8"))
    data = json.load(open(sys.argv[2], encoding="utf-8"))
    links = json.load(open(sys.argv[3], encoding="utf-8"))
    build(crawl, data["issues"], links, sys.argv[4])
